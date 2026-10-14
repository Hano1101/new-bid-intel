import io
from decimal import Decimal

import pytest
import xlwt
from docx import Document
from openpyxl import Workbook
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.cidfonts import UnicodeCIDFont
from reportlab.platypus import SimpleDocTemplate, Table, TableStyle

from app.bidder_result_details import (
    _package,
    extract_bidder_details,
    merge_bidder_details,
    normalize_bidder_name,
    parse_bidder_detail_rows,
)
from app.parsers import SourceDocument

ROWS = [
    ["供应商名称", "资格性审查", "符合性审查", "投标报价（万元）", "综合得分"],
    ["供应商甲", "通过", "通过", "12.3456", "99.92"],
    ["供应商乙", "不通过资格性审查，原因是：证明材料缺失", "", "", ""],
]


def _html(rows, before="", after=""):
    cells = ["<tr>" + "".join(f"<td>{cell}</td>" for cell in row) + "</tr>" for row in rows]
    return SourceDocument("notice.html", (before + "<table>" + "".join(cells) + "</table>" + after).encode())


def test_review_requires_both_checks_and_preserves_failed_reason():
    details = extract_bidder_details(_html(ROWS, "<p>合同包1</p>"))
    assert len(details) == 2
    passed, failed = details
    assert passed["package_code"] == "1"
    assert passed["qualification_status"] == passed["compliance_status"] == "通过"
    assert passed["review_status"] == "通过"
    assert passed["bid_amount"] == Decimal(123456)
    assert passed["evaluation_score"] == Decimal("99.92")
    assert failed["review_status"] == "未通过"
    assert "证明材料缺失" in failed["unsuccessful_reason"]
    assert failed["bid_amount"] is None
    assert "投标报价（万元）" in passed["field_evidence"]["bid_amount"]["source_evidence"]
    assert passed["source_location"] == "table:1/row:2"
    assert "供应商甲 通过 通过" in passed["scope_evidence"]


@pytest.mark.parametrize("headers,values,expected", [
    (["供应商", "资格审查"], ["公司甲", "通过"], None),
    (["供应商", "符合性审查"], ["公司甲", "通过"], None),
    (["供应商", "通过审查"], ["公司甲", "通过"], "通过"),
    (["供应商", "资格审查", "符合性审查"], ["公司甲", "通过", "未公布"], None),
    (["供应商", "资格审查", "符合性审查"], ["公司甲", "通过", "未通过"], "未通过"),
])
def test_review_is_never_inferred_from_partial_pass(headers, values, expected):
    details = extract_bidder_details(_html([headers, values]))
    assert details[0]["review_status"] == expected


@pytest.mark.parametrize("reason,expected", [
    ("资格性审查不通过，原因是材料缺失", "未通过"),
    ("未中标原因：得分较低", None),
])
def test_explicit_failed_review_reason_sets_review_status(reason, expected):
    detail = extract_bidder_details(_html([
        ["供应商", "未中标原因"], ["公司甲", reason],
    ]))[0]
    assert detail["review_status"] == expected
    if expected == "未通过":
        assert reason in detail["field_evidence"]["review_status"]["source_evidence"]


def test_final_quote_preferred_without_substituting_evaluation_or_award_amount():
    rows = [
        ["供应商", "投标报价（元）", "评审价格（元）", "中标金额（元）", "最终报价（元）"],
        ["公司甲", "50000", "38000", "42000", "41000"],
    ]
    assert extract_bidder_details(_html(rows))[0]["bid_amount"] == Decimal(41000)
    assert extract_bidder_details(_html([
        ["供应商", "资格审查", "中标金额（元）", "评审价格（元）"],
        ["公司甲", "通过", "42000", "38000"],
    ]))[0]["bid_amount"] is None


def test_unitless_bid_stays_empty_and_unrelated_wanyuan_description_does_not_anchor():
    rows = [["供应商", "最终报价"], ["公司甲", "120000"]]
    detail = extract_bidder_details(_html(rows, after="<p>代理服务费金额（万元）：1.20</p>"))[0]
    assert detail["bid_amount"] is None
    assert "单位" in detail["field_evidence"]["bid_amount"]["source_evidence"]


def test_same_table_explicit_amount_on_page_anchors_currency_for_all_bids():
    rows = [["供应商", "最终报价"], ["公司甲", "1,222,263.00"], ["公司乙", "1,212,228.00"]]
    details = extract_bidder_details(_html(rows, before="<p>公司甲成交金额：1,222,263.00元</p>"))
    assert [d["bid_amount"] for d in details] == [Decimal(1222263), Decimal(1212228)]
    assert "报价列单位" in details[1]["field_evidence"]["bid_amount"]["source_evidence"]


def test_conflicting_unit_anchors_are_not_assumed():
    rows = [["供应商", "投标报价"], ["公司甲", "12"], ["公司乙", "13"]]
    details = extract_bidder_details(_html(rows, before="<p>金额为12元；另为12万元；金额为13万元</p>"))
    # Unrelated amounts, even identical figures elsewhere on the page, cannot anchor.
    assert all(d["bid_amount"] is None for d in details)
    details = extract_bidder_details(_html(rows, before="<p>金额为12元；金额为13万元</p>"))
    assert all(d["bid_amount"] is None for d in details)


def test_context_currency_units_and_percent_quotes():
    details = extract_bidder_details(_html(
        [["供应商", "报价"], ["公司甲", "12.5"], ["公司乙", "95%"]],
        before="<p>报价单位：万元</p>",
    ))
    assert details[0]["bid_amount"] == Decimal(125000)
    assert details[1]["bid_amount"] is None


def test_complementary_merge_does_not_cross_packages_and_conflicts_survive_remerge():
    def parse(package, headers, row):
        return parse_bidder_detail_rows([headers, row], source_file="notice.html", table_index=1,
                                       context=f"合同包{package}")[0]
    a = parse(1, ["供应商", "资格审查"], ["公司甲", "通过"])
    b = parse(1, ["供应商", "符合性审查", "投标报价（元）"], ["公司甲", "通过", "100"])
    c = parse(2, ["供应商", "投标报价（元）"], ["公司甲", "200"])
    merged = merge_bidder_details([a, b, c])
    assert len(merged) == 2
    assert merged[0]["review_status"] == "通过"
    assert merged[0]["bid_amount"] == 100
    assert merged[1]["bid_amount"] == 200
    conflict = parse(1, ["供应商", "投标报价（元）"], ["公司甲", "101"])
    result = merge_bidder_details([*merged, conflict])
    assert result[0]["bid_amount"] is None
    assert result[0]["field_evidence"]["bid_amount"]["conflict"] is True
    assert merge_bidder_details([result[0], b])[0]["bid_amount"] is None


def test_default_package_only_falls_back_to_a_unique_named_package():
    rows = [["供应商", "综合得分"], ["公司甲", "99"]]
    default = parse_bidder_detail_rows(rows, source_file="n.html", table_index=1)[0]
    p1 = {**default, "package_code": "1"}
    p2 = {**default, "package_code": "2"}
    assert len(merge_bidder_details([default, p1])) == 1
    assert len(merge_bidder_details([default, p1, p2])) == 3
    other_default = {**default, "source_location": "table:2/row:3"}
    assert len(merge_bidder_details([default, other_default])) == 2
    assert normalize_bidder_name(" 公司（甲） \n") == "公司(甲)"


@pytest.mark.parametrize("context,expected", [
    ("A包", "A"), ("A包（采购内容）", "A"), ("1包", "1"),
    ("合同包2 评审结果", "2"), ("未分包项目评审结果", "default"),
])
def test_package_context_extracts_explicit_prefix_and_suffix_labels(context, expected):
    assert _package(context) == expected


@pytest.mark.parametrize("format", ["docx", "xlsx", "xls", "pdf"])
def test_real_document_table_formats(format):
    buffer = io.BytesIO()
    if format == "docx":
        doc = Document()
        doc.add_paragraph("合同包1")
        table = doc.add_table(rows=len(ROWS), cols=len(ROWS[0]))
        for r, row in enumerate(ROWS):
            for c, value in enumerate(row):
                table.cell(r, c).text = value
        doc.save(buffer)
    elif format == "xlsx":
        book = Workbook()
        sheet = book.active
        sheet.title = "合同包1"
        for row in ROWS:
            sheet.append(row)
        book.save(buffer)
        book.close()
    elif format == "xls":
        book = xlwt.Workbook()
        sheet = book.add_sheet("合同包1")
        for r, row in enumerate(ROWS):
            for c, value in enumerate(row):
                sheet.write(r, c, value)
        book.save(buffer)
    else:
        pdfmetrics.registerFont(UnicodeCIDFont("STSong-Light"))
        table = Table(ROWS[:2], colWidths=[100, 80, 80, 110, 80])
        table.setStyle(TableStyle([
            ("FONTNAME", (0, 0), (-1, -1), "STSong-Light"),
            ("GRID", (0, 0), (-1, -1), 0.5, "#000000"),
        ]))
        SimpleDocTemplate(buffer).build([table])
    details = extract_bidder_details(SourceDocument(f"review.{format}", buffer.getvalue()))
    assert details[0]["raw_name"] == "供应商甲"
    assert details[0]["review_status"] == "通过"
    assert details[0]["bid_amount"] == Decimal(123456)
    assert details[0]["evaluation_score"] == Decimal("99.92")
    if format == "pdf":
        assert details[0]["source_location"].startswith("page:1/")
    else:
        assert details[0]["package_code"] == "1"


def test_unsupported_or_malformed_document_is_safe():
    assert extract_bidder_details(SourceDocument("bad.xlsx", b"not an xlsx")) == []
    assert extract_bidder_details(SourceDocument("body.txt", "供应商甲 通过 100元".encode())) == []


def test_legacy_office_conversion_uses_shared_document_converter(monkeypatch):
    document = Document()
    table = document.add_table(rows=2, cols=5)
    for row_index, row in enumerate(ROWS[:2]):
        for column_index, value in enumerate(row):
            table.cell(row_index, column_index).text = value
    source = io.BytesIO()
    document.save(source)
    expected_docx = source.getvalue()
    from app import legacy_documents

    monkeypatch.setattr(legacy_documents, "convert_doc", lambda _content: expected_docx)
    details = extract_bidder_details(SourceDocument("review.doc", b"legacy-placeholder"))
    assert len(details) == 1
    assert details[0]["raw_name"] == "供应商甲"
    assert details[0]["evaluation_score"] == Decimal("99.92")
