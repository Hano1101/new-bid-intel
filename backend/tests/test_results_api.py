from decimal import Decimal
from io import BytesIO

from fastapi.testclient import TestClient
from openpyxl import load_workbook

from app.config import settings
from app.main import app
from app.schemas import ImportResult, ItemCandidate, NoticeMetadata, ParticipantCandidate
from app.storage import connect, save_import


def _import_sample(database):
    result = ImportResult(
        notice_id=0,
        source_files=["t20261013_12345678.html"],
        items_found=1,
        metadata=NoticeMetadata(
            project_name="结果表导出测试项目",
            procurement_unit="测试采购单位",
            announced_total_award=Decimal(12345),
        ),
        items=[
            ItemCandidate(
                product_name="测试设备",
                category="信息设备",
                brand="测试品牌",
                model="M-1",
                quantity=Decimal(2),
                unit_price=Decimal(12345),
                total_price=Decimal(24690),
                source_file="t20261013_12345678.html",
                source_location="table:1/row:2",
                source_evidence="测试原文证据",
            ),
        ],
        participants=[
            ParticipantCandidate(
                organization_name="测试中标供应商",
                outcome="winner",
                award_amount=Decimal(12345),
                source_file="t20261013_12345678.html",
                source_location="table:2/row:2",
                source_evidence="测试中标原文证据",
            ),
            ParticipantCandidate(
                organization_name="测试参与供应商",
                outcome="nonwinner",
                source_file="t20261013_12345678.html",
                source_location="table:2/row:3",
                source_evidence="测试参与原文证据",
            ),
        ],
    )
    return save_import(database, result)


def test_result_export_matches_template_and_keeps_unavailable_fields_blank(tmp_path, monkeypatch):
    database = tmp_path / "results.db"
    monkeypatch.setattr(settings, "database_path", str(database))
    imported = _import_sample(database)

    with TestClient(app) as client:
        response = client.post(
            "/api/v1/notices/export-results",
            json={"notice_ids": [imported.notice_id]},
        )
        missing = client.post(
            "/api/v1/notices/export-results", json={"notice_ids": [9999]}
        )

    assert response.status_code == 200
    assert response.headers["content-type"].startswith(
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    assert missing.status_code == 404

    workbook = load_workbook(BytesIO(response.content), data_only=True)
    assert workbook.sheetnames == ["总表", "标的表", "投标人表"]
    total = workbook["总表"]
    assert [total.cell(1, column).value for column in range(1, 9)] == [
        "序号", "URL", "UUID", "标段/合同包", "项目名称", "采购单位", "中标供应商", "中标金额",
    ]
    assert total["B1"].fill.fgColor.rgb == "FFD9D9D9"
    assert total["B3"].value is None
    assert total["C3"].value == "t20261013_12345678"
    assert total["E3"].value == "结果表导出测试项目"
    assert total["G3"].value == "测试中标供应商"
    assert total["H3"].value == 1.23
    assert total["H3"].number_format == "0.00"

    item = workbook["标的表"]
    assert item["D3"].value == "测试设备"
    assert item["H3"].value == 2
    assert item["H3"].number_format == "General"
    assert item["I3"].value == 1.2345
    assert item["J3"].value == 2.469
    assert item["I3"].number_format == "0.####"
    assert item["K3"].value is None

    bidders = workbook["投标人表"]
    assert bidders["D3"].value == "测试中标供应商"
    assert bidders["E3"].value == "中标人"
    assert bidders["E1"].fill.fgColor.rgb == "FF1F4E79"
    assert bidders["F3"].value is None
    assert bidders["G3"].value is None
    assert bidders["D4"].value == "测试参与供应商"
    assert bidders["E4"].value == "参与人"


def test_item_export_keeps_small_positive_prices(tmp_path, monkeypatch):
    database = tmp_path / "small-prices.db"
    monkeypatch.setattr(settings, "database_path", str(database))
    imported = _import_sample(database)
    with connect(database) as connection:
        connection.execute(
            "UPDATE procurement_items SET unit_price='2.5', total_price='6.5' WHERE notice_id=?",
            (imported.notice_id,),
        )
        connection.commit()
    with TestClient(app) as client:
        response = client.post(
            "/api/v1/notices/export-results", json={"notice_ids": [imported.notice_id]},
        )
    assert response.status_code == 200
    workbook = load_workbook(BytesIO(response.content), data_only=True)
    assert workbook["标的表"]["I3"].value == 0.00025
    assert workbook["标的表"]["J3"].value == 0.00065
    assert workbook["标的表"]["I3"].number_format == "0.#####"
    assert workbook["总表"]["H3"].value == 1.23


def test_export_does_not_replace_disclosed_total_with_inconsistent_product(tmp_path, monkeypatch):
    database = tmp_path / "conflicting-prices.db"
    monkeypatch.setattr(settings, "database_path", str(database))
    imported = _import_sample(database)
    with connect(database) as connection:
        connection.execute("UPDATE procurement_items SET total_price='20000' WHERE notice_id=?",
                           (imported.notice_id,))
        connection.commit()
    with TestClient(app) as client:
        response = client.post("/api/v1/notices/export-results",
                               json={"notice_ids": [imported.notice_id]})
    sheet = load_workbook(BytesIO(response.content), data_only=True)["标的表"]
    assert sheet["J3"].value == 2
    assert sheet["K3"].value is None


def test_export_limits_rows_to_selected_notices_and_active_dataset(tmp_path, monkeypatch):
    database = tmp_path / "results.db"
    monkeypatch.setattr(settings, "database_path", str(database))
    first = _import_sample(database)
    second = save_import(database, ImportResult(
        notice_id=0, source_files=["t20261013_98765432.html"], items_found=0, items=[],
        metadata=NoticeMetadata(project_name="另一个项目"),
    ))
    with TestClient(app) as client:
        response = client.post(
            "/api/v1/notices/export-results", json={"notice_ids": [first.notice_id, first.notice_id]},
        )
        dataset = client.post("/api/v1/datasets", json={"name": "空的数据集"}).json()
        wrong_dataset = client.post(
            "/api/v1/notices/export-results", json={"notice_ids": [first.notice_id]},
            headers={"X-Dataset-ID": dataset["id"]},
        )
        empty = client.post("/api/v1/notices/export-results", json={"notice_ids": []})

    assert response.status_code == 200
    assert wrong_dataset.status_code == 404
    assert empty.status_code == 422
    workbook = load_workbook(BytesIO(response.content), data_only=True)
    assert workbook["总表"].max_row == 3
    assert second.notice_id != first.notice_id
    assert workbook["总表"]["C3"].value == "t20261013_12345678"


def test_export_retains_packages_with_missing_amounts_without_repeating_project_total(tmp_path, monkeypatch):
    database = tmp_path / "results.db"
    monkeypatch.setattr(settings, "database_path", str(database))
    first = _import_sample(database)
    second = save_import(database, ImportResult(
        notice_id=0, source_files=["t20261013_98765432.html"], items_found=0, items=[],
        metadata=NoticeMetadata(project_name="多包未提取金额项目", announced_total_award=Decimal(99999)),
    ))
    with connect(database) as connection:
        for notice_id, names in [(first.notice_id, ["已中标包", "空白包二", "空白包三"]),
                                 (second.notice_id, ["空白包甲", "空白包乙"])]:
            project_id = connection.execute(
                "SELECT id FROM projects WHERE notice_id=?", (notice_id,),
            ).fetchone()["id"]
            connection.execute(
                "UPDATE packages SET package_code='1', package_name=?, package_award_total=NULL WHERE project_id=?",
                (names[0], project_id),
            )
            for index, name in enumerate(names[1:], start=2):
                connection.execute(
                    "INSERT INTO packages(project_id, package_code, package_name) VALUES(?,?,?)",
                    (project_id, str(index), name),
                )
        connection.execute("UPDATE notices SET project_name='=1+1' WHERE id=?", (first.notice_id,))
        connection.commit()

    with TestClient(app) as client:
        response = client.post(
            "/api/v1/notices/export-results", json={"notice_ids": [first.notice_id, second.notice_id]},
        )
    assert response.status_code == 200
    workbook = load_workbook(BytesIO(response.content))
    total = workbook["总表"]
    assert total.max_row == 7
    assert [total.cell(row, 4).value for row in range(3, 8)] == [
        "已中标包", "空白包二", "空白包三", "空白包甲", "空白包乙",
    ]
    assert total["H3"].value == 1.23
    assert all(total.cell(row, 8).value is None for row in range(4, 8))
    assert total["E3"].value == "=1+1"
    assert total["E3"].data_type == "s"


def test_export_includes_evidence_backed_review_bid_score_and_failure_reason(tmp_path, monkeypatch):
    database = tmp_path / "results.db"
    monkeypatch.setattr(settings, "database_path", str(database))
    imported = _import_sample(database)
    winner = {
        "package_code": "default",
        "raw_name": "测试中标供应商",
        "qualification_status": "通过",
        "compliance_status": "通过",
        "review_status": "通过",
        "bid_amount": Decimal(12345),
        "evaluation_score": Decimal("99.925"),
        "unsuccessful_reason": None,
        "source_file": "review.html",
        "source_location": "table:5/row:2",
        "source_evidence": "测试供应商 | 通过 | 报价 12345元 | 综合得分 99.925",
        "field_evidence": {},
    }
    failed = {
        **winner,
        "raw_name": "测试参与供应商",
        "qualification_status": "未通过",
        "compliance_status": None,
        "review_status": "未通过",
        "bid_amount": None,
        "evaluation_score": None,
        "unsuccessful_reason": "不通过资格性审查，原因是证明材料缺失",
        "source_evidence": "资格性审查 | 不通过资格性审查，原因是证明材料缺失",
        "source_location": "table:5/row:3",
    }
    from app.bidder_result_details import normalize_bidder_name
    from app.storage import save_bidder_result_details
    with connect(database) as connection:
        save_bidder_result_details(connection, imported.notice_id, [winner, failed])
        connection.commit()

    with TestClient(app) as client:
        response = client.post(
            "/api/v1/notices/export-results", json={"notice_ids": [imported.notice_id]},
        )
    assert response.status_code == 200
    bidders = load_workbook(BytesIO(response.content), data_only=True)["投标人表"]
    rows = {
        normalize_bidder_name(bidders.cell(row, 4).value): row
        for row in range(3, bidders.max_row + 1)
    }
    winner_row = rows[normalize_bidder_name("测试中标供应商")]
    assert bidders.cell(winner_row, 5).value == "中标人"
    assert bidders.cell(winner_row, 6).value == "通过"
    assert bidders.cell(winner_row, 7).value == 1.23
    assert bidders.cell(winner_row, 8).value == 99.93
    assert bidders.cell(winner_row, 9).value is None

    failed_row = rows[normalize_bidder_name("测试参与供应商")]
    assert bidders.cell(failed_row, 5).value == "其他"
    assert bidders.cell(failed_row, 6).value == "未通过"
    assert bidders.cell(failed_row, 7).value is None
    assert "证明材料缺失" in bidders.cell(failed_row, 9).value


def test_explicit_bidder_package_gets_a_matching_summary_parent_row(tmp_path, monkeypatch):
    database = tmp_path / "explicit-package-parent.db"
    monkeypatch.setattr(settings, "database_path", str(database))
    imported = _import_sample(database)
    from app.storage import save_bidder_result_details

    record = {
        "package_code": "1", "raw_name": "附表新增投标人",
        "qualification_status": "通过", "compliance_status": "通过",
        "review_status": "通过", "bid_amount": Decimal(125000),
        "evaluation_score": Decimal("88.50"), "unsuccessful_reason": None,
        "source_file": "review.pdf", "source_location": "page:1/table:1/row:2",
        "source_evidence": "1包 | 附表新增投标人 | 通过 | 12.5万元 | 88.50",
        "field_evidence": {},
    }
    with connect(database) as connection:
        connection.execute("UPDATE packages SET package_code='default', package_name=NULL")
        save_bidder_result_details(connection, imported.notice_id, [record])
        connection.commit()

    with TestClient(app) as client:
        response = client.post(
            "/api/v1/notices/export-results", json={"notice_ids": [imported.notice_id]},
        )
    workbook = load_workbook(BytesIO(response.content), data_only=True)
    total = workbook["总表"]
    bidders = workbook["投标人表"]
    total_keys = {(total.cell(row, 3).value, total.cell(row, 4).value)
                  for row in range(3, total.max_row + 1)}
    bidder_keys = {(bidders.cell(row, 2).value, bidders.cell(row, 3).value)
                   for row in range(3, bidders.max_row + 1)}
    assert response.status_code == 200
    assert ("t20261013_12345678", "1") in total_keys
    assert ("t20261013_12345678", "1") in bidder_keys
