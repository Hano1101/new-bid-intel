"""Evidence-backed bidder fields for result workbooks, separate from core extraction.

This module neither changes award outcomes nor treats an award amount or an
evaluation price as a bid. A numeric bid without defensible currency units is
left empty, with its original cell retained as evidence.
"""
from __future__ import annotations

import io
import re
import struct
import unicodedata
import zipfile
from decimal import Decimal, InvalidOperation
from pathlib import PurePosixPath
from typing import Any

from bs4 import BeautifulSoup
from docx import Document
from openpyxl import load_workbook

from app.file_formats import inspect_content
from app.parsers import (
    PARTICIPANT_ORG_ALIASES,
    SourceDocument,
    _clean,
    _header_key,
    _invalid_participant_name,
    _money,
    _participant_package_code,
)

_FIELDS = (
    "qualification_status", "compliance_status", "review_status", "bid_amount",
    "evaluation_score", "unsuccessful_reason",
)
_ALIASES = {
    "raw_name": PARTICIPANT_ORG_ALIASES,
    "package_code": ("包号", "合同包", "合同包号", "采购包", "采购包号", "标段号", "标包"),
    "qualification_status": ("资格性审查", "资格审查", "资格性审查结果", "资格审查结果"),
    "compliance_status": ("符合性审查", "符合性检查", "符合性审查结果"),
    "review_status": ("通过审查", "审查结果", "审查结论", "综合审查结果", "综合审查"),
    "bid_amount": (
        "最终报价", "最终投标报价", "最终响应报价", "投标报价", "响应报价",
        "投标总报价", "响应总报价", "报价金额", "报价", "投标报价总额",
    ),
    "evaluation_score": ("综合得分", "评审总得分", "评审得分", "总得分", "得分"),
    "unsuccessful_reason": ("未中标原因", "未成交原因", "未中标成交原因", "落标原因"),
}
_PACKAGE_PATTERN = re.compile(
    r"(?:合同|采购|标段|标包)?包\s*(?:编号|号)?\s*[:：]?\s*"
    r"(?P<after>[A-Za-z0-9一二三四五六七八九十]+)(?![A-Za-z0-9])(?=\s*(?:[（(:：]|$|[\u4e00-\u9fff]))"
    r"|(?P<before>[A-Za-z0-9一二三四五六七八九十]+)\s*(?:号)?包"
)
_AMOUNT_PATTERN = re.compile(r"(-?\d[\d,，]*(?:\.\d+)?)\s*(亿元|万元|元)")
_UNIT_PATTERN = re.compile(r"(?:金额单位|报价单位|单位)\s*[：:]?\s*[（(]?\s*(?:人民币)?\s*(亿元|万元|元)")


def normalize_bidder_name(value: Any) -> str:
    """Normalize formatting only; never fuzzy-merge different branches/entities."""
    return re.sub(r"\s+", "", unicodedata.normalize("NFKC", _clean(value))).casefold()


def _status(value: str) -> str | None:
    key = _header_key(value)
    if any(token in key for token in ("不通过", "未通过", "不合格", "不符合", "无效")):
        return "未通过"
    if key in {"通过", "合格", "符合", "符合要求", "是", "审查通过", "资格性审查通过",
               "资格审查通过", "符合性审查通过", "综合审查通过", "综合通过", "评审通过",
               "审查合格", "全部通过", "均通过", "均合格"}:
        return "通过"
    return None


def _decimal(value: str) -> Decimal | None:
    value = re.sub(r"[\s,，]", "", unicodedata.normalize("NFKC", value))
    if not re.fullmatch(r"-?\d+(?:\.\d+)?", value):
        return None
    try:
        result = Decimal(value)
        return result if result.is_finite() else None
    except InvalidOperation:
        return None


def _unit(value: str) -> str | None:
    match = re.search(r"亿元|万元|元", value)
    return match.group(0) if match else None


def _columns(header: list[str], context: str) -> dict[str, int]:
    keys = [_header_key(cell) for cell in header]
    result = {}
    for field, aliases in _ALIASES.items():
        for alias in aliases:
            if _header_key(alias) in keys:
                result[field] = keys.index(_header_key(alias))
                break
    if ("raw_name" not in result and "名称" in keys
            and ({"qualification_status", "compliance_status", "review_status"} & result.keys()
                 or re.search(r"投标|供应商|响应", context))):
        result["raw_name"] = keys.index("名称")
    return result


def _package(context: str) -> str:
    matches = list(_PACKAGE_PATTERN.finditer(unicodedata.normalize("NFKC", context)))
    if not matches:
        return "default"
    match = matches[-1]
    return _participant_package_code(match.group("after") or match.group("before"))


def _nearby_result_package(table: Any) -> str:
    node = table
    for _ in range(8):
        parent = node.parent
        if parent is None:
            break
        preceding = []
        for child in parent.children:
            if child is node:
                break
            if getattr(child, "name", None) == "table":
                continue
            text = _clean(child if isinstance(child, str) else child.get_text(" ", strip=True))
            if text and len(text) <= 180:
                preceding.append(text)
        for text in reversed(preceding):
            if (code := _package(text)) != "default":
                return code
        node = parent
    return "default"


def _supplier_amount_units(name: str, raw_amount: str, source_text: str,
                          peer_names: list[str]) -> list[tuple[str, str]]:
    amount = _decimal(raw_amount)
    if amount is None:
        return []
    normalized_text = unicodedata.normalize("NFKC", source_text).casefold()
    normalized_name = unicodedata.normalize("NFKC", name).casefold()
    if not normalized_name:
        return []
    name_pattern = r"\s*".join(re.escape(char) for char in normalized_name)
    name_spans = [match.span() for match in re.finditer(name_pattern, normalized_text)]
    if not name_spans:
        return []
    peer_patterns = [r"\s*".join(re.escape(char) for char in unicodedata.normalize("NFKC", peer).casefold())
                     for peer in peer_names
                     if normalize_bidder_name(peer) and normalize_bidder_name(peer) != normalize_bidder_name(name)]
    matches = []
    for match in _AMOUNT_PATTERN.finditer(normalized_text):
        if _decimal(match.group(1)) != amount:
            continue
        for start, end in name_spans:
            # The explicit currency evidence must follow this supplier name in
            # nearby source text; a coincidentally equal earlier amount is not
            # enough to attribute the unit to this bidder.
            distance = match.start() - end
            if distance < 0 or distance > 180:
                continue
            between = normalized_text[end:match.start()]
            if any(re.search(peer, between) for peer in peer_patterns):
                continue
            matches.append((match.group(2), match.group(0)))
    return matches


def _table_unit(header: str, context: str) -> tuple[str | None, str]:
    if unit := _unit(header):
        return unit, header
    matches = list(_UNIT_PATTERN.finditer(context))
    units = {match.group(1) for match in matches}
    if len(units) == 1:
        return units.pop(), " / ".join(match.group(0) for match in matches)
    return None, "金额未标单位，且未找到同一供应商同额的显式原文金额"


def parse_bidder_detail_rows(
    rows: list[list[Any]], *, source_file: str, table_index: int,
    context: str = "", source_text: str = "", location_prefix: str = "",
) -> list[dict[str, Any]]:
    """Read explicit bidder/review columns with table-local package evidence."""
    rows = [[_clean(cell) for cell in row] for row in rows]
    records = []
    header_indices = [
        index for index, row in enumerate(rows)
        if "raw_name" in (columns := _columns(row, context))
        and any(field in columns for field in _FIELDS)
    ]
    for block, header_index in enumerate(header_indices):
        end = header_indices[block + 1] if block + 1 < len(header_indices) else len(rows)
        header = rows[header_index]
        local_context = " ".join([context, *(" ".join(row) for row in rows[max(0, header_index - 4):header_index])])
        columns = _columns(header, local_context)
        package_code = _package(local_context)
        unit, unit_evidence = (None, "")
        if "bid_amount" in columns:
            if "raw_name" not in columns:
                continue
            unit, unit_evidence = _table_unit(header[columns["bid_amount"]], local_context)
            amount_column = columns["bid_amount"]
            peers = [row[columns["raw_name"]] for row in rows[header_index + 1:end]
                     if columns["raw_name"] < len(row)]
            row_unit_evidence = {
                row_index: _supplier_amount_units(
                    rows[row_index][columns["raw_name"]], rows[row_index][amount_column],
                    source_text, peers,
                )
                for row_index in range(header_index + 1, end)
                if columns["raw_name"] < len(rows[row_index])
                and amount_column < len(rows[row_index])
            }
            # Procurement tables use one unit for an entire quotation column.
            # A same-supplier, same-value match to an explicit amount elsewhere
            # in the notice establishes that shared unit; unrelated page amounts
            # never do. If the source contains conflicting anchors, propagate none.
            evidenced_units = {
                anchor[0] for anchors in row_unit_evidence.values() for anchor in anchors
            }
            if unit is not None:
                evidenced_units.add(unit)
            if unit is None and len(evidenced_units) == 1:
                unit = next(iter(evidenced_units))
                anchor_supplier = next(
                    rows[index][columns["raw_name"]]
                    for index, anchors in row_unit_evidence.items() if anchors
                )
                unit_evidence = (
                    f"同表 {anchor_supplier} 的报价与原文同额且明确单位，"
                    f"据此确认该报价列单位为{unit}"
                )
        for row_index in range(header_index + 1, end):
            row = rows[row_index]

            def cell(field, current_row=row, current_columns=columns):
                column = current_columns.get(field)
                return current_row[column] if column is not None and column < len(current_row) else ""

            raw_name = cell("raw_name")
            if _invalid_participant_name(raw_name) or _header_key(raw_name) in {
                _header_key(alias) for alias in (*PARTICIPANT_ORG_ALIASES, "名称")
            }:
                continue
            if cell("package_code"):
                package_code = _participant_package_code(cell("package_code"))
            location = f"{location_prefix}table:{table_index}/row:{row_index + 1}"
            record = dict.fromkeys(_FIELDS)
            record.update(
                package_code=package_code, raw_name=raw_name,
                source_file=source_file, source_location=location, field_evidence={},
                source_evidence="\n".join(filter(None, [local_context, " | ".join(header), " | ".join(row)])),
                scope_evidence=" ".join(row),
            )
            for field in _FIELDS:
                raw = cell(field)
                if not raw:
                    continue
                evidence = {
                    "source_file": source_file,
                    "source_location": f"{location}/column:{columns[field] + 1}",
                    "source_evidence": f"{header[columns[field]]}\n{raw}",
                }
                record["field_evidence"][field] = evidence
                if field.endswith("status"):
                    record[field] = _status(raw)
                elif field == "bid_amount":
                    own_unit = _unit(raw) or unit
                    own_unit_evidence = raw if _unit(raw) else unit_evidence
                    record[field] = _money(raw, own_unit or "") if own_unit else None
                    evidence["source_evidence"] += f"\n单位依据：{own_unit_evidence}"
                elif field == "evaluation_score":
                    record[field] = _decimal(raw)
                else:
                    record[field] = raw
            failed = [
                cell(field) for field in ("qualification_status", "compliance_status", "review_status")
                if record[field] == "未通过"
            ]
            if failed:
                record["review_status"] = "未通过"
                reason = "；".join(dict.fromkeys(filter(None, [record["unsuccessful_reason"], *failed])))
                record["unsuccessful_reason"] = reason
                record["field_evidence"]["unsuccessful_reason"] = {
                    "source_file": source_file, "source_location": location, "source_evidence": reason,
                }
            elif (record["unsuccessful_reason"] and re.search(
                    r"(?:不通过|未通过|不合格|不符合|无效)",
                    unicodedata.normalize("NFKC", record["unsuccessful_reason"]),
                ) and re.search(
                    r"资格|符合性|审查|评审",
                    unicodedata.normalize("NFKC", record["unsuccessful_reason"]),
                )):
                record["review_status"] = "未通过"
                record["field_evidence"]["review_status"] = dict(
                    record["field_evidence"].get("unsuccessful_reason", {}),
                )
            elif record["qualification_status"] == record["compliance_status"] == "通过":
                record["review_status"] = "通过"
            if record["review_status"] is not None and "review_status" not in record["field_evidence"]:
                record["field_evidence"]["review_status"] = {
                    "source_file": source_file, "source_location": location,
                    "source_evidence": "\n".join(cell(field) for field in ("qualification_status", "compliance_status")),
                }
            records.append(record)
    return records


def extract_bidder_details(document: SourceDocument) -> list[dict[str, Any]]:
    """Extract sidecar fields from supported tables; malformed/unavailable tables yield no rows."""
    detected = inspect_content(document.content)
    if detected.error:
        return []
    suffix = "." + detected.format if detected.format else PurePosixPath(document.filename).suffix.lower()
    records = []

    def add(rows, index, context="", source_text="", prefix=""):
        records.extend(parse_bidder_detail_rows(
            rows, source_file=document.filename, table_index=index, context=context,
            source_text=source_text, location_prefix=prefix,
        ))

    try:
        if suffix in {".doc", ".rtf", ".wps"}:
            from app.legacy_documents import convert_doc

            return extract_bidder_details(SourceDocument(document.filename, convert_doc(document.content)))
        if suffix in {".html", ".htm"}:
            soup = BeautifulSoup(document.content, "html.parser")
            source_text = soup.get_text(" ", strip=True)
            for index, table in enumerate(soup.find_all("table"), 1):
                rows = [
                    ["" if cell.find("table") else cell.get_text(" ", strip=True)
                     for cell in row.find_all(["th", "td"], recursive=False)]
                    for row in table.find_all("tr") if row.find_parent("table") is table
                ]
                context = []
                sibling = table.previous_sibling
                while sibling is not None and len(context) < 4:
                    if isinstance(sibling, str):
                        if sibling.strip():
                            context.append(_clean(sibling))
                    elif getattr(sibling, "name", None) in {"p", "h1", "h2", "h3", "h4", "strong", "b", "br"}:
                        text = sibling.get_text(" ", strip=True)
                        if text and len(text) <= 180:
                            context.append(text)
                    else:
                        break
                    sibling = sibling.previous_sibling
                if (package := _nearby_result_package(table)) != "default":
                    context.append(f"合同包{package}")
                add(rows, index, " ".join(reversed(context)), source_text)
        elif suffix == ".docx":
            word = Document(io.BytesIO(document.content))
            source_text = " ".join([*(p.text for p in word.paragraphs),
                                    *(cell.text for table in word.tables for row in table.rows for cell in row.cells)])
            from docx.oxml.ns import qn
            from docx.table import Table
            from docx.text.paragraph import Paragraph
            context = []
            index = 0
            for child in word.element.body.iterchildren():
                if child.tag == qn("w:p"):
                    text = Paragraph(child, word).text
                    if text.strip():
                        context.append(text)
                elif child.tag == qn("w:tbl"):
                    index += 1
                    table = Table(child, word)
                    add([[cell.text for cell in row.cells] for row in table.rows], index,
                        " ".join(context[-4:]), source_text)
                    context = []
        elif suffix == ".xlsx":
            workbook = load_workbook(io.BytesIO(document.content), read_only=True, data_only=True)
            try:
                for index, sheet in enumerate(workbook.worksheets, 1):
                    rows = [[_clean(cell) for cell in row] for row in sheet.iter_rows(values_only=True)]
                    add(rows, index, sheet.title, " ".join(" ".join(row) for row in rows))
            finally:
                workbook.close()
        elif suffix == ".xls":
            import xlrd
            workbook = xlrd.open_workbook(file_contents=document.content, on_demand=True)
            try:
                for index, sheet in enumerate(workbook.sheets(), 1):
                    rows = [[_clean(cell.value) for cell in row] for row in sheet.get_rows()]
                    add(rows, index, sheet.name, " ".join(" ".join(row) for row in rows))
            finally:
                workbook.release_resources()
        elif suffix == ".pdf":
            import pdfplumber
            with pdfplumber.open(io.BytesIO(document.content)) as pdf:
                for page_number, page in enumerate(pdf.pages, 1):
                    source_text = page.extract_text() or ""
                    for index, table in enumerate(page.find_tables(), 1):
                        _, top, _, _ = table.bbox
                        context = page.crop((0, max(0, top - 70), page.width, top)).extract_text() or ""
                        add(table.extract(), index, context, source_text, f"page:{page_number}/")
    except (ValueError, TypeError, OSError, KeyError, IndexError, ImportError,
            zipfile.BadZipFile, struct.error):
        # Sidecar failure must not prevent the existing import from completing.
        return records
    return records


def merge_bidder_details(records: list[dict[str, Any]]) -> list[dict[str, Any]]:
    """Merge complementary evidence; conflicting fields stay empty, including on remerge."""
    named_packages: dict[str, set[str]] = {}
    for record in records:
        name = normalize_bidder_name(record.get("raw_name"))
        package = _participant_package_code(record.get("package_code"))
        if package != "default":
            named_packages.setdefault(name, set()).add(package)
    groups: dict[tuple[str, str, str], list[dict[str, Any]]] = {}
    for index, record in enumerate(records):
        name = normalize_bidder_name(record.get("raw_name"))
        package = _participant_package_code(record.get("package_code"))
        if package == "default" and len(named_packages.get(name, set())) == 1:
            package = next(iter(named_packages[name]))
        # An unresolved package is not proof that two records belong together.
        # Retain each such source row until an unambiguous named package anchors it.
        source_key = ((record.get("source_file") or "") + "\0"
                      + (record.get("source_location") or str(index))) if package == "default" else ""
        groups.setdefault((package, name, source_key), []).append(record)
    output = []
    for (package, _, _), rows in groups.items():
        merged = {**rows[0], "package_code": package, "field_evidence": {}}
        merged["source_evidence"] = "\n".join(dict.fromkeys(
            row.get("source_evidence", "") for row in rows if row.get("source_evidence")
        ))
        conflicts = set()
        for field in _FIELDS:
            values = []
            evidences = []
            for row in rows:
                value = row.get(field)
                if value is not None and value != "" and value not in values:
                    values.append(value)
                evidence = row.get("field_evidence", {}).get(field)
                if evidence:
                    evidences.append(evidence)
                    if evidence.get("conflict"):
                        conflicts.add(field)
            if len(values) > 1:
                conflicts.add(field)
            merged[field] = values[0] if len(values) == 1 and field not in conflicts else None
            if evidences or field in conflicts:
                merged["field_evidence"][field] = {
                    "source_file": " / ".join(dict.fromkeys(e.get("source_file", "") for e in evidences)),
                    "source_location": " / ".join(dict.fromkeys(e.get("source_location", "") for e in evidences)),
                    "source_evidence": "\n".join(dict.fromkeys(e.get("source_evidence", "") for e in evidences)),
                }
                if field in conflicts:
                    merged["field_evidence"][field].update(conflict=True, values=[str(v) for v in values])
        if "review_status" not in conflicts and not ({"qualification_status", "compliance_status"} & conflicts):
            if "未通过" in (merged["qualification_status"], merged["compliance_status"]):
                merged["review_status"] = "未通过"
            elif merged["qualification_status"] == merged["compliance_status"] == "通过":
                merged["review_status"] = "通过"
        output.append(merged)
    return output
